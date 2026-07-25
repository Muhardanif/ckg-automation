"""
Cek cepat logika pembacaan & normalisasi data peserta (app/readers.py).

Ini bagian yang paling mahal kalau salah: data yang meleset di sini terkirim
ke portal dan baru ketahuan setelah ditolak Dukcapil, satu per satu.

JALANKAN:  venv\\Scripts\\python.exe test_readers.py
Tanpa framework — cukup assert. Keluar kode 0 = semua lolos.
"""
import os
import tempfile
from datetime import date, datetime

import openpyxl

from app.readers import (
    _bersihkan, _teks, _format_tanggal, _normalisasi_jk, _tgl_dari_nik,
    baca_excel, validasi, cek_konsistensi_nik,
    koreksi_tgl_dari_nik, koreksi_jk_dari_nik,
)
from app.excel_hasil import cari_kolom, simpan_workbook, tgl_iso
from app.schema import Peserta, KelompokUsia

# NIK contoh: ...DDMMYY... digit 7-12. Perempuan = DD + 40.
NIK_P = "3515064804640001"   # perempuan, lahir 08-04-1964
NIK_L = "3515060804640002"   # laki-laki, lahir 08-04-1964


def _peserta(**kw):
    dasar = dict(nik=NIK_L, nama="Budi", tgl_lahir="1964-04-08",
                 jenis_kelamin="L", kelompok_usia=KelompokUsia.DEWASA)
    dasar.update(kw)
    return Peserta(**dasar)


def test_bersihkan():
    assert _bersihkan(None) is None
    assert _bersihkan("  Budi  ") == "Budi"
    # sel berisi spasi saja = kosong, supaya nilai default tetap dipakai
    assert _bersihkan("   ") is None
    # angka bulat selalu jadi teks rapi, apa pun tipe aslinya dari openpyxl
    assert _bersihkan(70.0) == "70"
    assert _bersihkan(70) == "70"
    assert _bersihkan(70.5) == 70.5       # pecahan dibiarkan angka
    assert _bersihkan(True) is True       # TRUE/FALSE bukan "1"/"0"


def test_teks():
    # sel NIK/telepon bertipe angka di Excel harus keluar sebagai string
    assert _teks(8123456789) == "8123456789"
    assert _teks(8123456789.0) == "8123456789"
    assert _teks("08123456789") == "08123456789"   # 0 di depan dipertahankan
    assert _teks(None) is None
    assert _teks("  ") is None


def test_format_tanggal():
    assert _format_tanggal(None) is None
    assert _format_tanggal(datetime(1964, 4, 8, 13, 30)) == "1964-04-08"
    assert _format_tanggal(date(1964, 4, 8)) == "1964-04-08"
    assert _format_tanggal("1964-04-08") == "1964-04-08"
    assert _format_tanggal("08/04/1964") == "1964-04-08"   # hari/bulan, bukan US
    assert _format_tanggal("08-04-1964") == "1964-04-08"
    assert _format_tanggal("  1964-04-08  ") == "1964-04-08"
    # nomor seri Excel (sel tanggal yang terlanjur jadi angka)
    assert _format_tanggal(23475) == "1964-04-08"
    assert _format_tanggal("23475") == "1964-04-08"
    # tidak dikenali -> diteruskan apa adanya, ditandai oleh validasi()
    assert _format_tanggal("bukan tanggal") == "bukan tanggal"


def test_normalisasi_jk():
    for v in ("L", "l", "Laki-laki", "PRIA", "male", " M "):
        assert _normalisasi_jk(v) == "L", v
    for v in ("P", "Perempuan", "wanita", "FEMALE", " f "):
        assert _normalisasi_jk(v) == "P", v
    assert _normalisasi_jk(None) is None


def test_tgl_dari_nik():
    assert _tgl_dari_nik(NIK_L) == (8, 4, 64, "L")
    assert _tgl_dari_nik(NIK_P) == (8, 4, 64, "P")
    assert _tgl_dari_nik("123") is None                  # panjang salah
    assert _tgl_dari_nik("35150648046400A1") is None     # ada huruf
    assert _tgl_dari_nik("3515069904640001") is None     # DD=99 mustahil


def test_validasi():
    assert validasi(_peserta()) == []
    assert "NIK harus 16 digit angka" in validasi(_peserta(nik="123"))
    assert "Nama kosong" in validasi(_peserta(nama=""))
    assert "Tanggal lahir kosong" in validasi(_peserta(tgl_lahir=""))
    assert "Jenis kelamin tidak valid" in validasi(_peserta(jenis_kelamin="X"))


def test_cek_konsistensi_nik():
    assert cek_konsistensi_nik(_peserta()) == []
    # tanggal beda dgn NIK -> 1 peringatan
    assert len(cek_konsistensi_nik(_peserta(tgl_lahir="1964-08-04"))) == 1
    # jenis kelamin beda dgn NIK -> 1 peringatan
    assert len(cek_konsistensi_nik(_peserta(nik=NIK_P))) == 1


def test_koreksi_tgl_dari_nik():
    # sudah cocok -> tidak perlu dikoreksi
    assert koreksi_tgl_dari_nik(_peserta()) is None
    # hari/bulan tertukar, 2 digit tahun sama -> abad dari Excel dipertahankan
    assert koreksi_tgl_dari_nik(_peserta(tgl_lahir="1964-08-04")) == "1964-04-08"
    # 2 digit tahun beda -> abad ditebak; 2064 di masa depan, jadi 1964
    assert koreksi_tgl_dari_nik(_peserta(tgl_lahir="1970-01-01")) == "1964-04-08"
    # tanggal kosong -> tetap bisa diturunkan dari NIK
    assert koreksi_tgl_dari_nik(_peserta(tgl_lahir="")) == "1964-04-08"
    # NIK tak valid -> jangan mengarang
    assert koreksi_tgl_dari_nik(_peserta(nik="123")) is None
    # 31 Februari (DD=31, MM=02) mustahil -> jangan koreksi
    assert koreksi_tgl_dari_nik(_peserta(nik="3515063102640001",
                                         tgl_lahir="1970-01-01")) is None


def test_koreksi_jk_dari_nik():
    assert koreksi_jk_dari_nik(_peserta()) is None           # sudah cocok
    assert koreksi_jk_dari_nik(_peserta(nik=NIK_P)) == "P"   # NIK bilang P
    assert koreksi_jk_dari_nik(_peserta(nik="123")) is None


def test_baca_excel():
    """Regresi pembacaan Excel: header tidak di baris 1, sel tanggal asli,
    NIK bertipe angka, sel kosong, dan baris kosong di tengah."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Laporan CKG"])                    # baris judul (bukan header)
    ws.append(["NIK", "Nama", "Tanggal Lahir", "Jenis Kelamin", "No. HP",
               "Alamat", "Pekerjaan", "BB (kg)", "Tekanan Darah"])
    ws.append([NIK_L, "Budi", date(1964, 4, 8), "Laki-laki", "08123456789",
               "Jl. Mawar", "Petani", 70.0, "120/80"])
    ws.append([None, None, None, None, None, None, None, None, None])  # kosong
    ws.append([int(NIK_P), "Siti", "08/04/1964", "P", 8129999999,
               None, None, 55, None])
    path = os.path.join(tempfile.mkdtemp(), "uji.xlsx")
    wb.save(path)

    hasil = baca_excel(path, KelompokUsia.DEWASA, header_row=1)
    assert len(hasil) == 2, "baris kosong harus dilewati"

    a, b = hasil
    assert a.nik == NIK_L
    assert a.nama == "Budi"
    assert a.tgl_lahir == "1964-04-08"          # sel bertipe tanggal
    assert a.jenis_kelamin == "L"
    assert a.no_hp == "08123456789"             # 0 di depan tidak hilang
    assert a.pemeriksaan["berat_badan"] == "70"
    assert a.pemeriksaan["tekanan_darah"] == "120/80"
    assert a.pekerjaan == "Petani"
    # header di baris ke-2 (0-based 1), data pertama di baris ke-3 Excel
    assert a.baris_sumber == 3
    assert a.file_sumber == "uji.xlsx"

    assert b.nik == NIK_P                       # NIK bertipe angka -> string utuh
    assert b.tgl_lahir == "1964-04-08"          # teks dd/mm/yyyy
    assert b.jenis_kelamin == "P"
    assert b.no_hp == "8129999999"
    assert b.alamat is None
    assert b.pekerjaan == "Lainnya"             # sel kosong -> nilai default
    assert b.pemeriksaan["kolesterol"] is None  # kolom tak ada di file
    assert b.baris_sumber == 5

    # sheet tanpa baris apa pun tidak boleh melempar error
    wb2 = openpyxl.Workbook()
    path2 = os.path.join(tempfile.mkdtemp(), "kosong.xlsx")
    wb2.save(path2)
    assert baca_excel(path2, KelompokUsia.DEWASA, header_row=5) == []


def test_tgl_iso():
    assert tgl_iso(None) is None
    assert tgl_iso(datetime(2026, 6, 12, 10, 30)) == "2026-06-12"
    assert tgl_iso("2026-06-12T10:30:00") == "2026-06-12"
    assert tgl_iso("2026-06-12 10:30") == "2026-06-12"
    assert tgl_iso("bukan tanggal") is None


def test_cari_kolom_dan_simpan():
    """Helper bersama yang dipakai ketiga tool tahap (dulu disalin 3x)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["NIK", " Nama ", "No. Tiket"])

    assert cari_kolom(ws, 1, "NIK") == 1
    assert cari_kolom(ws, 1, "Nama") == 2          # header ber-spasi tetap cocok
    assert cari_kolom(ws, 1, "Status Daftar") is None
    # buat=True menambah kolom baru di ujung, lalu ketemu di pencarian berikutnya
    baru = cari_kolom(ws, 1, "Status Daftar", buat=True)
    assert baru == 4
    assert cari_kolom(ws, 1, "Status Daftar") == 4
    assert cari_kolom(ws, 1, "Status Daftar", buat=True) == 4   # tidak dobel

    path = os.path.join(tempfile.mkdtemp(), "simpan.xlsx")
    assert simpan_workbook(wb, path, "UJI") is True
    assert os.path.exists(path)
    # tak bisa ditulis -> False, bukan exception (kasus nyata: dibuka di Excel)
    assert simpan_workbook(wb, tempfile.mkdtemp(), "UJI") is False


if __name__ == "__main__":
    for nama, fn in sorted(globals().items()):
        if nama.startswith("test_"):
            fn()
            print(f"OK  {nama}")
    print("\nSemua cek lolos.")
